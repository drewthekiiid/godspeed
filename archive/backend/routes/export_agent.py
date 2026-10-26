from flask import Blueprint, request, send_file, jsonify
import io
from docx import Document

export_bp = Blueprint("export", __name__)

@export_bp.route("/export", methods=["POST"])
def export_estimate():
    try:
        data = request.get_json()
        estimate = data.get("estimate", [])
        if not estimate:
            return jsonify({"error": "Missing estimate data"}), 400

        doc = Document()
        doc.add_heading("Estimate Summary", 0)
        table = doc.add_table(rows=1, cols=6)
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'Division'
        hdr_cells[1].text = 'Scope Item'
        hdr_cells[2].text = 'Qty'
        hdr_cells[3].text = 'Unit'
        hdr_cells[4].text = 'Unit Price'
        hdr_cells[5].text = 'Total'

        for item in estimate:
            row_cells = table.add_row().cells
            row_cells[0].text = item.get("division", "")
            row_cells[1].text = item.get("scope_item", "")
            row_cells[2].text = str(item.get("quantity", ""))
            row_cells[3].text = item.get("unit", "")
            row_cells[4].text = f"${item.get('unit_price', 0):,.2f}"
            row_cells[5].text = f"${item.get('total', 0):,.2f}"

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)

        return send_file(buffer, as_attachment=True, download_name="estimate.docx", mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    except Exception as e:
        return jsonify({"error": f"Export failed: {str(e)}"}), 500